"""Tests for document-bounded identity propagation and span-accurate redaction."""

from __future__ import annotations

import concurrent.futures
from pathlib import Path
import tempfile
import unittest

import openpyxl
import docx
import fitz

from redaction_tool.detector import Match, detect, detect_summary, CATEGORY_MAP
from redaction_tool.ledger import EntityLedger
from redaction_tool.ocr import OCRWord, detect_on_ocr
from redaction_tool.redactor import ScanOptions, redact_file, extract_text


class TestDocumentContext(unittest.TestCase):

    def test_ms_patel_alias_from_credible_full_name_seed(self) -> None:
        text = "Patient Priya Patel was seen on Monday. Ms. Patel requested a copy."
        matches = detect(text, enabled_categories=["names"])
        matched_texts = {m.text for m in matches}
        self.assertIn("Priya Patel", matched_texts)
        self.assertIn("Ms. Patel", matched_texts)
        # Verify ledger evidence was attached to the alias
        patel_matches = [m for m in matches if m.text == "Ms. Patel"]
        self.assertTrue(any("ledger" in m.evidence for m in patel_matches))

    def test_brown_and_may_not_mass_redacted(self) -> None:
        text = "In May 2025, Alexandria Montoya filed a report in the brown folder."
        matches = detect(text, enabled_categories=["names", "dates"])
        matched_texts = {m.text for m in matches}
        self.assertNotIn("brown", matched_texts)
        # "May 2025" should be matched as dates, not name
        date_matches = [m for m in matches if m.category == "dates"]
        self.assertTrue(any("May 2025" in m.text for m in date_matches))
        name_matches = [m for m in matches if m.category == "names"]
        self.assertTrue(all("May" not in m.text for m in name_matches))

    def test_two_independent_documents_no_cross_leakage(self) -> None:
        doc1_text = "Alexandria Montoya was present at the hearing."
        doc2_text = "The file for Montoya."

        # Scan Document 1
        matches1 = detect(doc1_text, enabled_categories=["names"])
        self.assertTrue(any(m.text == "Alexandria Montoya" for m in matches1))

        # Scan Document 2 with a clean invocation (no ledger passed)
        matches2 = detect(doc2_text, enabled_categories=["names"])
        # Montoya in Doc 2 should NOT match as a ledger alias because Doc 2 has no seed
        self.assertFalse(any(m.text == "Montoya" for m in matches2))

    def test_concurrency_isolation(self) -> None:
        def scan_doc_a() -> list[str]:
            res = detect("Alexandria Montoya visited the clinic. Montoya signed.", enabled_categories=["names"])
            return [m.text for m in res]

        def scan_doc_b() -> list[str]:
            res = detect("Priya Patel was admitted. Patel was discharged.", enabled_categories=["names"])
            return [m.text for m in res]

        with concurrent.futures.ThreadPoolExecutor(max_workers=4) as executor:
            futures_a = [executor.submit(scan_doc_a) for _ in range(10)]
            futures_b = [executor.submit(scan_doc_b) for _ in range(10)]

            for f in futures_a:
                res = f.result()
                self.assertIn("Alexandria Montoya", res)
                self.assertNotIn("Priya Patel", res)
            for f in futures_b:
                res = f.result()
                self.assertIn("Priya Patel", res)
                self.assertNotIn("Alexandria Montoya", res)

    def test_explicit_disabled_names(self) -> None:
        text = "Dr. Jane Smith visited patient Alexandria Montoya. Email: jane@example.org"
        matches = detect(text, enabled_categories=["emails", "ssn"])
        categories = {m.category for m in matches}
        self.assertNotIn("names", categories)
        self.assertIn("emails", categories)
        self.assertEqual(len([m for m in matches if m.category == "names"]), 0)

    def test_use_presidio_false_behavior_unchanged(self) -> None:
        text = "Dr. Jane Smith visited the clinic on 05/12/2024."
        matches = detect(text, enabled_categories=["names", "dates"], use_presidio=False)
        matched_texts = {m.text for m in matches}
        self.assertIn("Dr. Jane Smith", matched_texts)

    def test_synthetic_cross_page_pdf_name_and_surname(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            pdf_path = Path(td) / "test_crosspage.pdf"
            doc = fitz.open()
            p1 = doc.new_page()
            p1.insert_text((50, 50), "Patient Alexandria Montoya was admitted to the hospital.")
            p2 = doc.new_page()
            p2.insert_text((50, 50), "Montoya reported feeling better after treatment.")
            doc.save(str(pdf_path))
            doc.close()

            out_path = Path(td) / "test_crosspage_REDACTED.pdf"
            opts = ScanOptions(enabled_categories=["names"])
            result = redact_file(pdf_path, opts, out_path=out_path)

            self.assertGreaterEqual(result.redaction_count, 2)
            redacted_doc = fitz.open(str(out_path))
            p1_text = redacted_doc[0].get_text()
            p2_text = redacted_doc[1].get_text()
            redacted_doc.close()

            self.assertNotIn("Alexandria Montoya", p1_text)
            self.assertNotIn("Montoya", p2_text)

    def test_docx_full_name_and_titled_surname(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            docx_path = Path(td) / "test_doc.docx"
            doc = docx.Document()
            doc.add_paragraph("Dr. Alex Smith reviewed the medical report.")
            doc.add_paragraph("Some intervening text about hospital records.")
            doc.add_paragraph("Ms. Smith signed off on the document.")
            doc.save(str(docx_path))

            out_path = Path(td) / "test_doc_REDACTED.docx"
            opts = ScanOptions(enabled_categories=["names"])
            result = redact_file(docx_path, opts, out_path=out_path)

            self.assertGreaterEqual(result.redaction_count, 2)
            redacted_doc = docx.Document(str(out_path))
            p0_text = redacted_doc.paragraphs[0].text
            p2_text = redacted_doc.paragraphs[2].text

            self.assertNotIn("Alex Smith", p0_text)
            self.assertNotIn("Ms. Smith", p2_text)

    def test_xlsx_name_and_later_alias(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            xlsx_path = Path(td) / "test_sheet.xlsx"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws["A1"] = "Alexandria Montoya"
            ws["C10"] = "Montoya"
            wb.save(str(xlsx_path))
            wb.close()

            out_path = Path(td) / "test_sheet_REDACTED.xlsx"
            opts = ScanOptions(enabled_categories=["names"])
            result = redact_file(xlsx_path, opts, out_path=out_path)

            self.assertGreaterEqual(result.redaction_count, 2)
            redacted_wb = openpyxl.load_workbook(str(out_path))
            redacted_ws = redacted_wb.active
            self.assertEqual(redacted_ws["A1"].value, "[REDACTED]")
            self.assertEqual(redacted_ws["C10"].value, "[REDACTED]")
            redacted_wb.close()

    def test_enabled_categories_empty_list_xlsx(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            xlsx_path = Path(td) / "test_empty_cats.xlsx"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws["A1"] = "First Name"
            ws["A2"] = "Jane Smith"
            ws["B1"] = "Email"
            ws["B2"] = "jane@example.com"
            wb.save(str(xlsx_path))
            wb.close()

            out_path = Path(td) / "test_empty_cats_REDACTED.xlsx"
            opts = ScanOptions(enabled_categories=[])
            result = redact_file(xlsx_path, opts, out_path=out_path)

            self.assertEqual(result.redaction_count, 0)
            redacted_wb = openpyxl.load_workbook(str(out_path))
            redacted_ws = redacted_wb.active
            self.assertEqual(redacted_ws["A2"].value, "Jane Smith")
            self.assertEqual(redacted_ws["B2"].value, "jane@example.com")
            redacted_wb.close()

    def test_punctuation_heavy_and_unicode_cell_span_replacement(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            xlsx_path = Path(td) / "test_unicode.xlsx"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws["A1"] = "Contact: (555) 019-2831, e-mail: test@example.org & Dr. José-María O'Brien"
            wb.save(str(xlsx_path))
            wb.close()

            out_path = Path(td) / "test_unicode_REDACTED.xlsx"
            opts = ScanOptions(enabled_categories=["phones", "emails", "names"])
            result = redact_file(xlsx_path, opts, out_path=out_path)

            self.assertGreaterEqual(result.redaction_count, 3)
            redacted_wb = openpyxl.load_workbook(str(out_path))
            redacted_ws = redacted_wb.active
            val = redacted_ws["A1"].value
            self.assertNotIn("(555) 019-2831", val)
            self.assertNotIn("test@example.org", val)
            self.assertNotIn("José-María", val)
            self.assertIn("[REDACTED]", val)
            redacted_wb.close()

    def test_ocr_unit_name_then_alias_with_bboxes(self) -> None:
        words_unit1 = [
            OCRWord("Alexandria", 10, 10, 80, 20, 90.0, 0, 10),
            OCRWord("Montoya", 100, 10, 70, 20, 95.0, 11, 18),
        ]
        words_unit2 = [
            OCRWord("Patient", 10, 50, 50, 20, 88.0, 0, 7),
            OCRWord("Montoya", 70, 50, 70, 20, 92.0, 8, 15),
            OCRWord("spoke", 150, 50, 40, 20, 85.0, 16, 21),
        ]

        doc_ledger = EntityLedger()
        opts = ScanOptions(enabled_categories=["names"])

        entries1, _ = detect_on_ocr(words_unit1, scan_opts=opts, ledger=doc_ledger)
        self.assertTrue(len(entries1) >= 1)

        entries2, _ = detect_on_ocr(words_unit2, scan_opts=opts, ledger=doc_ledger)
        self.assertTrue(len(entries2) >= 1)
        montoya_entry = entries2[0]
        self.assertEqual(montoya_entry["entity_type"], "names")
        self.assertEqual(montoya_entry["bbox"], [70.0, 50.0, 140.0, 70.0])

    def test_scan_redact_equivalent_category_coverage(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            docx_path = Path(td) / "test_parity.docx"
            doc = docx.Document()
            doc.add_paragraph("Dr. Jane Smith (SSN: 123-45-6789) emailed jane@example.com on 01/15/2024.")
            doc.save(str(docx_path))

            text = extract_text(docx_path)
            opts = ScanOptions(enabled_categories=["names", "ssn", "emails", "dates"])
            scan_matches = opts.run_detection(text)

            scan_counts: dict[str, int] = {}
            for m in scan_matches:
                scan_counts[m.category] = scan_counts.get(m.category, 0) + 1

            out_path = Path(td) / "test_parity_REDACTED.docx"
            result = redact_file(docx_path, opts, out_path=out_path)

            self.assertEqual(result.redaction_count, len(scan_matches))
            for cat, count in scan_counts.items():
                self.assertEqual(result.per_category.get(cat, 0), count)

    def test_overlap_union_property(self) -> None:
        text = "Patient 123-45-6789 (DOB: 01/02/1990)"
        matches = detect(text, enabled_categories=["unique_ids", "ssn", "dates"])
        # Verify both SSN and unique_ids/dates are detected and covered without losing character spans
        categories = {m.category for m in matches}
        self.assertIn("ssn", categories)
        self.assertIn("dates", categories)


if __name__ == "__main__":
    unittest.main()
