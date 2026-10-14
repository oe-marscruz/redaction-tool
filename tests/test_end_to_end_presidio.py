"""End-to-end integration tests for Presidio-assisted redaction, category mapping,

and offline security verification.
"""

from __future__ import annotations

import hashlib
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import MagicMock, patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from redaction_tool import detector, redactor, verify


class TestEndToEndPresidio(unittest.TestCase):
    """End-to-end verification of Presidio detection, category isolation, and export security."""

    def setUp(self):
        self.tmp_dir = tempfile.TemporaryDirectory()
        self.td = Path(self.tmp_dir.name)

    def tearDown(self):
        self.tmp_dir.cleanup()

    def test_use_presidio_false_end_to_end_pdf_docx_xlsx(self):
        """use_presidio=False must complete end-to-end redactions without touching Presidio."""
        try:
            import pymupdf as fitz
        except ImportError:
            import fitz
        import docx
        import openpyxl

        # 1. PDF
        pdf_src = self.td / "sample.pdf"
        doc = fitz.open()
        p = doc.new_page()
        p.insert_text((50, 50), "Student Brianna Ynostroza SSN 123-45-6789 jane@example.edu")
        doc.save(str(pdf_src))
        doc.close()

        # 2. DOCX
        docx_src = self.td / "sample.docx"
        d = docx.Document()
        d.add_paragraph("Contact Brianna Ynostroza at jane@example.edu or 123-45-6789.")
        d.save(str(docx_src))

        # 3. XLSX
        xlsx_src = self.td / "sample.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "Brianna Ynostroza"
        ws["A2"] = "123-45-6789"
        wb.save(str(xlsx_src))
        wb.close()

        opts = redactor.ScanOptions(use_presidio=False)

        with patch.object(detector, "get_presidio_engine", side_effect=AssertionError("Presidio called when disabled")):
            for src in (pdf_src, docx_src, xlsx_src):
                out = self.td / f"{src.stem}_REDACTED{src.suffix}"
                res = redactor.redact_file(src, opts, out_path=out)
                self.assertIsNone(res.error)
                self.assertGreater(res.redaction_count, 0)
                v_res = verify.verify_export(out, targeted_values=["123-45-6789", "jane@example.edu"])
                self.assertEqual(v_res["status"], "PASS")

    def test_fail_closed_when_presidio_errors_in_end_to_end(self):
        """When use_presidio=True and Presidio engine raises, file redaction must fail closed."""
        try:
            import pymupdf as fitz
        except ImportError:
            import fitz

        pdf_src = self.td / "fail_closed.pdf"
        doc = fitz.open()
        p = doc.new_page()
        p.insert_text((50, 50), "Patient Zyrphonia was seen in clinic.")
        doc.save(str(pdf_src))
        doc.close()

        opts = redactor.ScanOptions(use_presidio=True)

        class BrokenPresidio:
            def analyze(self, **kwargs):
                raise RuntimeError("Presidio engine internal error")

        with patch.object(detector, "get_presidio_engine", return_value=BrokenPresidio()):
            out = self.td / "fail_closed_REDACTED.pdf"
            res = redactor.redact_file(pdf_src, opts, out_path=out)
            self.assertIsNotNone(res.error)
            self.assertIn("Presidio analysis failed", str(res.error))

    def test_category_whitelist_isolation_end_to_end(self):
        """With use_presidio=True and enabled_categories=['names'], only names are requested/redacted."""
        if not detector.presidio_is_available():
            self.skipTest("Presidio not installed in environment")

        try:
            import pymupdf as fitz
        except ImportError:
            import fitz

        pdf_src = self.td / "category_test.pdf"
        doc = fitz.open()
        p = doc.new_page()
        p.insert_text((50, 50), "Patient Alice Johnson lives in Boston, SSN 214-55-1234.")
        doc.save(str(pdf_src))
        doc.close()

        out = self.td / "category_test_REDACTED.pdf"
        opts = redactor.ScanOptions(enabled_categories=["names"], use_presidio=True)
        res = redactor.redact_file(pdf_src, opts, out_path=out)
        self.assertIsNone(res.error)

        redacted_text = redactor.extract_text(out)
        self.assertNotIn("Alice Johnson", redacted_text)
        self.assertIn("214-55-1234", redacted_text)

    def test_offline_security_guarantee_presidio(self):
        """Presidio detection and redaction operate entirely offline with zero socket calls."""
        if not detector.presidio_is_available():
            self.skipTest("Presidio not installed in environment")

        def blocked_socket(*args, **kwargs):
            raise RuntimeError("Network socket call detected during offline execution")

        text = "Contact Alice Johnson at Boston office or email alice@example.com."
        opts = redactor.ScanOptions(enabled_categories=["names", "addresses", "emails"], use_presidio=True)

        with patch("socket.socket", side_effect=blocked_socket):
            matches = opts.run_detection(text)
            self.assertGreater(len(matches), 0)

    def test_presidio_ledger_propagation_end_to_end(self):
        """Presidio-detected full name seeds document ledger for subsequent surname-only redactions."""
        try:
            import pymupdf as fitz
        except ImportError:
            import fitz

        pdf_src = self.td / "ledger_test.pdf"
        doc = fitz.open()
        p1 = doc.new_page()
        p1.insert_text((50, 50), "Patient Zyrphonia Silva was admitted to the ward.")
        p2 = doc.new_page()
        p2.insert_text((50, 50), "Silva showed signs of improvement on Tuesday.")
        doc.save(str(pdf_src))
        doc.close()

        fake_match = detector.Match(
            category="names",
            start=8,
            end=23,
            text="Zyrphonia Silva",
            confidence="confirmed",
            evidence="presidio:PERSON:0.95",
        )

        out = self.td / "ledger_test_REDACTED.pdf"
        opts = redactor.ScanOptions(enabled_categories=["names"], use_presidio=True)

        with patch.object(detector, "presidio_is_available", return_value=True):
            with patch.object(detector, "_presidio_matches", return_value=[fake_match]):
                res = redactor.redact_file(pdf_src, opts, out_path=out)
                self.assertIsNone(res.error)
                self.assertGreaterEqual(res.redaction_count, 2)

                doc_red = fitz.open(str(out))
                t1 = doc_red[0].get_text()
                t2 = doc_red[1].get_text()
                doc_red.close()

                self.assertNotIn("Zyrphonia Silva", t1)
                self.assertNotIn("Silva", t2)


if __name__ == "__main__":
    unittest.main()
